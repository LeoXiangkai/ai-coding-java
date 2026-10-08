#!/usr/bin/env python3
"""
Generate a styled Excel (.xlsx) test case file from JSON input.

Supports hierarchical module grouping (up to 5 levels) with merged cells.

Usage:
    python generate_xlsx.py --input test_cases.json --output test_cases.xlsx
    python generate_xlsx.py --output test_cases.xlsx              # reads stdin
    python generate_xlsx.py --input tc.json --output tc.xlsx --lang en
    python generate_xlsx.py --input tc.json --output tc.xlsx --max-group-level 4
    python generate_xlsx.py --input tc.json --output tc.xlsx --extra-columns "Assignee,Sprint"

Requires: openpyxl (install it in the project environment; this component does not install dependencies)
"""

import argparse
import json
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import (
        Alignment, Border, Font, PatternFill, Side,
    )
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl is required; install it in the project environment or use Markdown/CSV output.", file=sys.stderr)
    sys.exit(2)


# ── Column definitions ──────────────────────────────────────────────────

GROUP_COLUMNS_ZH = [
    ("group_l1", "一级分组", 20),
    ("group_l2", "二级分组", 20),
    ("group_l3", "三级分组", 18),
    ("group_l4", "四级分组", 16),
    ("group_l5", "五级分组", 16),
]

GROUP_COLUMNS_EN = [
    ("group_l1", "Group L1", 20),
    ("group_l2", "Group L2", 20),
    ("group_l3", "Group L3", 18),
    ("group_l4", "Group L4", 16),
    ("group_l5", "Group L5", 16),
]

DATA_COLUMNS_ZH = [
    ("id",              "编号",           12),
    ("test_item",       "测试项",         30),
    ("priority",        "优先级",         10),
    ("requirement_id",  "需求ID",         16),
    ("title",           "用例标题",       40),
    ("preconditions",   "前置条件",       30),
    ("steps",           "操作步骤",       50),
    ("expected",        "预期结果",       40),
    ("test_data",       "测试数据",       30),
    ("is_reverse",      "是否反向用例",   14),
    ("type",            "测试类型",       14),
    ("automation_type", "关联自动化类型", 18),
    ("ai_generated",    "AI生成",         12),
]

DATA_COLUMNS_EN = [
    ("id",              "ID",              14),
    ("test_item",       "Test Item",       30),
    ("priority",        "Priority",        10),
    ("requirement_id",  "Requirement ID",  16),
    ("title",           "Title",           40),
    ("preconditions",   "Preconditions",   30),
    ("steps",           "Steps",           50),
    ("expected",        "Expected Result", 40),
    ("test_data",       "Test Data",       30),
    ("is_reverse",      "Is Reverse Case", 14),
    ("type",            "Test Type",       14),
    ("automation_type", "Automation Type", 18),
    ("ai_generated",    "AI Generated",    12),
]

# ── Styles ───────────────────────────────────────────────────────────────

FONT_NAME = "Microsoft YaHei"

HEADER_FONT = Font(name=FONT_NAME, bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

GROUP_HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")

BODY_FONT = Font(name=FONT_NAME, size=10)
BODY_ALIGN = Alignment(vertical="top", wrap_text=True)
BODY_CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

GROUP_FONT = Font(name=FONT_NAME, size=10, bold=True, color="2F5496")
GROUP_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14, color="4472C4")
TITLE_ALIGN = Alignment(horizontal="left", vertical="center")

THIN_BORDER = Border(
    left=Side(style="thin", color="B4C6E7"),
    right=Side(style="thin", color="B4C6E7"),
    top=Side(style="thin", color="B4C6E7"),
    bottom=Side(style="thin", color="B4C6E7"),
)

STRIPE_FILL = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")

GROUP_FILLS = [
    PatternFill(start_color="E9EFF7", end_color="E9EFF7", fill_type="solid"),  # L1
    PatternFill(start_color="EDF2FA", end_color="EDF2FA", fill_type="solid"),  # L2
    PatternFill(start_color="F2F6FC", end_color="F2F6FC", fill_type="solid"),  # L3
    PatternFill(start_color="F7F9FD", end_color="F7F9FD", fill_type="solid"),  # L4
    PatternFill(start_color="FAFBFE", end_color="FAFBFE", fill_type="solid"),  # L5
]

PRIORITY_FILLS = {
    "P1": PatternFill(start_color="FF4444", end_color="FF4444", fill_type="solid"),
    "P2": PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid"),
    "P3": PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid"),
    "P4": PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid"),
}
PRIORITY_FONTS = {
    "P1": Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF"),
    "P2": Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF"),
    "P3": Font(name=FONT_NAME, size=10, bold=True),
    "P4": Font(name=FONT_NAME, size=10),
}

CENTER_FIELDS = {"id", "priority", "type", "requirement_id", "is_reverse",
                 "automation_type", "ai_generated"}


# ── Merge helper ─────────────────────────────────────────────────────────

def _merge_group_columns(ws, header_row, test_cases, columns, group_count):
    """Merge adjacent cells with same value in each group column."""
    if group_count == 0 or not test_cases:
        return

    first_data_row = header_row + 1
    total_rows = len(test_cases)

    for g_idx in range(group_count):
        field = columns[g_idx][0]
        col_num = g_idx + 1

        i = 0
        while i < total_rows:
            current_val = test_cases[i].get(field, "")
            # Also require parent groups to match for merging
            j = i + 1
            while j < total_rows:
                if test_cases[j].get(field, "") != current_val:
                    break
                # Check all parent group levels also match
                parent_match = True
                for p in range(g_idx):
                    pf = columns[p][0]
                    if test_cases[j].get(pf, "") != test_cases[i].get(pf, ""):
                        parent_match = False
                        break
                if not parent_match:
                    break
                j += 1

            if j - i > 1:
                ws.merge_cells(
                    start_row=first_data_row + i,
                    start_column=col_num,
                    end_row=first_data_row + j - 1,
                    end_column=col_num,
                )
            i = j


# ── Main generator ───────────────────────────────────────────────────────

def generate_xlsx(data, output_path, lang="zh", max_group_level=3,
                  extra_columns=None):
    """Generate a styled Excel workbook from test case data."""
    max_group_level = max(0, min(5, max_group_level))

    group_cols_all = GROUP_COLUMNS_ZH if lang == "zh" else GROUP_COLUMNS_EN
    data_cols = list(DATA_COLUMNS_ZH if lang == "zh" else DATA_COLUMNS_EN)

    group_cols = list(group_cols_all[:max_group_level])
    columns = group_cols + data_cols

    if extra_columns:
        for col_name in extra_columns:
            columns.append((col_name.lower().replace(" ", "_"), col_name, 18))

    wb = Workbook()
    ws = wb.active
    ws.title = data.get("project", "Test Cases")[:31]  # sheet name max 31 chars

    # ── Title row ────────────────────────────────────────────────────────
    project = data.get("project", "")
    title_text = project or "Test Cases"

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    title_cell = ws.cell(row=1, column=1, value=title_text)
    title_cell.font = TITLE_FONT
    title_cell.alignment = TITLE_ALIGN
    ws.row_dimensions[1].height = 36

    # ── Header row ───────────────────────────────────────────────────────
    header_row = 3
    for col_idx, (field, header, width) in enumerate(columns, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
        # Group columns use darker header
        if field.startswith("group_l"):
            cell.fill = GROUP_HEADER_FILL
        else:
            cell.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[header_row].height = 28

    # ── Data rows ────────────────────────────────────────────────────────
    test_cases = data.get("test_cases", [])
    for row_offset, tc in enumerate(test_cases):
        row_num = header_row + 1 + row_offset
        is_odd = row_offset % 2 == 1

        for col_idx, (field, _, _) in enumerate(columns, 1):
            value = tc.get(field, "")
            cell = ws.cell(row=row_num, column=col_idx, value=value)
            cell.border = THIN_BORDER

            # Group columns: special styling
            if field.startswith("group_l"):
                cell.font = GROUP_FONT
                cell.alignment = GROUP_ALIGN
                level = int(field[-1]) - 1
                if level < len(GROUP_FILLS):
                    cell.fill = GROUP_FILLS[level]
            elif field in CENTER_FIELDS:
                cell.alignment = BODY_CENTER
                # Priority color coding
                if field == "priority" and value in PRIORITY_FILLS:
                    cell.fill = PRIORITY_FILLS[value]
                    cell.font = PRIORITY_FONTS[value]
                elif is_odd:
                    cell.fill = STRIPE_FILL
                    cell.font = BODY_FONT
                else:
                    cell.font = BODY_FONT
            else:
                cell.alignment = BODY_ALIGN
                if is_odd:
                    cell.fill = STRIPE_FILL
                cell.font = BODY_FONT

        # Auto row height based on content
        max_lines = 1
        for col_idx, (field, _, _) in enumerate(columns, 1):
            value = str(tc.get(field, ""))
            lines = value.count("\n") + 1
            max_lines = max(max_lines, lines)
        ws.row_dimensions[row_num].height = max(20, min(max_lines * 16, 200))

    # ── Merge group columns ──────────────────────────────────────────────
    _merge_group_columns(ws, header_row, test_cases, group_cols, max_group_level)

    # ── Summary row ──────────────────────────────────────────────────────
    if test_cases:
        summary_row = header_row + 1 + len(test_cases) + 1
        total_label = "Total" if lang == "en" else "合计"

        # Count by type
        type_counts = {}
        for tc in test_cases:
            t = tc.get("type", "Other")
            type_counts[t] = type_counts.get(t, 0) + 1
        type_summary = ", ".join(
            "{}:{}".format(k, v) for k, v in sorted(type_counts.items())
        )

        ws.merge_cells(
            start_row=summary_row, start_column=1,
            end_row=summary_row, end_column=max_group_level or 1,
        )
        ws.cell(row=summary_row, column=1,
                value="{}: {}".format(total_label, len(test_cases)))
        ws.cell(row=summary_row, column=1).font = Font(
            name=FONT_NAME, bold=True, size=10, color="4472C4"
        )

        type_col = max_group_level + 11  # type column index
        if type_col <= len(columns):
            ws.cell(row=summary_row, column=max_group_level + 1,
                    value=type_summary)
            ws.cell(row=summary_row, column=max_group_level + 1).font = Font(
                name=FONT_NAME, size=9, color="666666"
            )

    # ── Freeze panes & auto-filter ───────────────────────────────────────
    # Freeze after group columns + header row
    freeze_col = get_column_letter(max_group_level + 1) if max_group_level else "A"
    ws.freeze_panes = "{}{}".format(freeze_col, header_row + 1)
    ws.auto_filter.ref = "A{}:{}{}".format(
        header_row,
        get_column_letter(len(columns)),
        header_row + len(test_cases),
    )

    # ── Save ─────────────────────────────────────────────────────────────
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output))
    print("OK: {} test cases ({} groups) written to {}".format(
        len(test_cases), max_group_level, output
    ))
    return str(output)


def _default_output_path(data):
    """Build a portable default output path in the current project directory."""
    project = data.get("project", "TestCases")
    # Sanitize project name for filename
    safe_name = "".join(c if c.isalnum() or c in ("-", "_", " ") else "_"
                        for c in project).strip() or "TestCases"
    return str(Path.cwd() / "{}_测试用例.xlsx".format(safe_name))


def main():
    parser = argparse.ArgumentParser(description="Generate test case Excel file")
    parser.add_argument("--input", "-i", help="JSON input file (default: stdin)")
    parser.add_argument("--output", "-o", default="",
                        help="Output .xlsx path (default: ./<project>_测试用例.xlsx)")
    parser.add_argument("--lang", default="zh", choices=["zh", "en"],
                        help="Column header language (default: zh)")
    parser.add_argument("--max-group-level", type=int, default=3,
                        help="Number of group columns 1-5 (default: 3)")
    parser.add_argument("--extra-columns", default="",
                        help="Comma-separated extra column names")
    args = parser.parse_args()

    if args.input:
        with open(args.input, "r", encoding="utf-8") as f:
            data = json.load(f)
    else:
        data = json.load(sys.stdin)

    output_path = args.output if args.output else _default_output_path(data)

    extra = [c.strip() for c in args.extra_columns.split(",") if c.strip()] \
        if args.extra_columns else None

    generate_xlsx(data, output_path, lang=args.lang,
                  max_group_level=args.max_group_level,
                  extra_columns=extra)


if __name__ == "__main__":
    main()
